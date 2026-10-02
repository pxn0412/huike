import io
import tempfile
import subprocess
import unittest
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement
from PIL import Image

import parsing
import chunking
from store import Store


def word_bytes():
    doc = Document()
    doc.add_heading('报名要求', level=1)
    doc.add_paragraph('每队人数为1至5人。')
    table = doc.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [('材料', '要求'), ('演示视频', '10分钟以内')]):
        for cell, text in zip(row.cells, values):
            cell.text = text
    doc.add_paragraph('表格之后的截止日期为2026年9月24日。')
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def image_bytes(format='PNG'):
    out = io.BytesIO()
    Image.new('RGB', (80, 40), 'white').save(out, format)
    return out.getvalue()


class FormatTests(unittest.TestCase):
    def test_doc_timeout_requests_owned_word_cleanup(self):
        with patch('document_formats.subprocess.run', side_effect=[subprocess.TimeoutExpired('word', 90), None]) as run:
            with self.assertRaises(ValueError):
                parsing.extract_document('old.doc', bytes.fromhex('D0CF11E0A1B11AE1'))
        self.assertEqual(run.call_count, 2)
        self.assertIn('-CleanupOnly', run.call_args.args[0])

    def test_word_content_controls_are_not_silently_skipped(self):
        doc = Document()
        doc.add_paragraph('Start')
        sdt = OxmlElement('w:sdt'); body = OxmlElement('w:sdtContent')
        paragraph = doc.add_paragraph('Deadline: October 1, 2026.')
        body.append(paragraph._p); sdt.append(body)
        doc.element.body.insert(1, sdt)
        table = doc.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        inner = OxmlElement('w:sdt'); content = OxmlElement('w:sdtContent')
        item = cell.add_paragraph('Team size: 1-5')
        content.append(item._p); inner.append(content); cell._tc.append(inner)
        raw = io.BytesIO(); doc.save(raw)
        pages, _ = parsing.extract_document('form.docx', raw.getvalue(), False)
        self.assertIn('Deadline: October 1, 2026.', pages[0])
        self.assertIn('Team size: 1-5', pages[0])

    def test_image_ocr_result_becomes_searchable_text(self):
        with patch('document_vision.configured', return_value=False), patch('document_formats.os.name', 'nt'), patch('document_formats.run_image_ocr',
                return_value=[{'file': 'page-0001.png', 'text': '报名人数为1至5人。'}]):
            pages, meta = parsing.extract_document('notice.png', image_bytes())
        self.assertEqual(pages, ['报名人数为1至5人。'])
        self.assertEqual(meta['ocr_pages'], [1])

    def test_empty_and_oversized_upload_rejected(self):
        for content in (b'', b'x' * (20 * 1024 * 1024 + 1)):
            with self.assertRaises(ValueError):
                parsing.extract_document('x.txt', content)

    def test_docx_keeps_paragraph_table_order_and_no_fake_pages(self):
        pages, meta = parsing.extract_document('notice.docx', word_bytes(), use_ocr=False)
        text = '\n'.join(pages)
        self.assertLess(text.index('1至5人'), text.index('演示视频'))
        self.assertLess(text.index('10分钟以内'), text.index('表格之后'))
        self.assertIn('| 材料 | 要求 |', text)
        self.assertEqual(meta['location_kind'], 'section')
        self.assertEqual(meta['source_type'], 'docx')

    def test_images_are_validated_and_unavailable_ocr_is_visible(self):
        for ext, format in [('png', 'PNG'), ('jpg', 'JPEG'), ('webp', 'WEBP'), ('bmp', 'BMP'), ('tiff', 'TIFF')]:
            with self.subTest(ext=ext):
                pages, meta = parsing.extract_document('notice.' + ext, image_bytes(format), use_ocr=False)
                self.assertEqual(meta['location_kind'], 'image')
                self.assertEqual(meta['unreadable_pages'], [1])
                self.assertTrue(meta['warning'])
                self.assertEqual(pages, [''])

    def test_wrong_extension_and_corrupt_content_rejected(self):
        for filename, content in [('x.exe', b'abc'), ('x.png', b'not png'),
                                  ('x.docx', b'PK-not-word'), ('x.pdf', image_bytes()),
                                  ('x.jpg', image_bytes()), ('x.doc', b'not word')]:
            with self.subTest(filename=filename):
                with self.assertRaises(ValueError):
                    parsing.extract_document(filename, content, use_ocr=False)

    def test_text_utf8_and_gb18030(self):
        for encoding in ('utf-8-sig', 'gb18030'):
            pages, meta = parsing.extract_document('通知.txt', '报名截止日期：九月。'.encode(encoding))
            self.assertIn('报名截止', pages[0])
            self.assertEqual(meta['location_kind'], 'section')


class MultiFormatStoreTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.store = Store(self.tmp.name, use_ocr=False)
        for target, value in [('ai.configured', False), ('store.clients.agent_client.configured', False)]:
            p = patch(target, return_value=value)
            p.start()
            self.addCleanup(p.stop)
        p = patch('store.clients.knowledge_base.push', return_value={'status': 'uploaded', 'doc_id': 'test'})
        self.push = p.start()
        self.addCleanup(p.stop)

    def test_word_upload_confirm_search_preview_and_duplicate(self):
        raw = word_bytes()
        upload = self.store.upload('通知.docx', raw)
        data = dict(upload_id=upload['upload_id'], title='通知', name='测试赛', edition='2026', uploader='测试')
        result = self.store.confirm(data)
        did = result['document_id']
        self.assertEqual(self.store.file('documents', did), raw)
        preview = self.store.preview('documents', did)
        self.assertEqual(preview['location_kind'], 'section')
        self.assertIn('演示视频', '\n'.join(preview['pages']))
        hits = self.store.search(result['competition_id'], '演示视频')
        self.assertTrue(hits)
        self.assertEqual(hits[0]['location_kind'], 'section')
        remote_text = self.push.call_args[0][0].decode()
        self.assertIn('正文', remote_text)
        self.assertNotIn('第 1 页', remote_text)
        self.assertNotIn('两赛道通用', remote_text)
        repeat = self.store.upload('不同文件名.docx', raw)
        self.assertTrue(self.store.confirm({**data, 'upload_id': repeat['upload_id'],
                                            'competition_id': result['competition_id']})['duplicate'])

    def test_cross_page_chunk_does_not_claim_single_page(self):
        upload = self.store.upload('note.txt', b'Unique quotation sufficiently long for source matching.')
        result = self.store.confirm(dict(upload_id=upload['upload_id'], title='Note', name='Test', edition='2026', uploader='Test'))
        with self.store.db() as db:
            db.execute('UPDATE chunks SET page=1,page_end=2 WHERE document_id=?', (result['document_id'],))
            self.assertIsNone(self.store._unique_quote_page(db, result['document_id'],
                              'Unique quotation sufficiently long'))


class StructureChunkTests(unittest.TestCase):
    def test_large_heading_not_removed_when_split(self):
        heading = '# ' + '长标题' * 10
        blocks = chunking.build_blocks([(1, heading + '\n' + '甲' * 300)], max_chars=64)
        self.assertIn(heading, '\n'.join(b['text'] for b in blocks))

    def test_long_section_is_bounded_and_keeps_every_sentence(self):
        sentences = [f'第{i}项要求：' + '必须按照比赛通知提交作品' * 3 + '。' for i in range(40)]
        blocks = chunking.build_blocks([(1, '一、作品要求\n' + ''.join(sentences))], max_chars=240)
        self.assertGreater(len(blocks), 1)
        self.assertTrue(all(len(b['text']) <= 240 for b in blocks))
        self.assertTrue(all('作品要求' in b['heading'] for b in blocks))
        for sentence in sentences:
            self.assertTrue(any(sentence in b['text'] for b in blocks), sentence)

    def test_table_splits_by_rows_and_repeats_header(self):
        header = '| 材料名称 | 提交要求 |\n| --- | --- |'
        rows = [f'| 材料{i} | ' + '完整作品演示要求' * 4 + ' |' for i in range(20)]
        blocks = chunking.build_blocks([(1, '# 提交材料\n' + header + '\n' + '\n'.join(rows))], max_chars=220)
        self.assertGreater(len(blocks), 1)
        for block in blocks:
            self.assertIn('| 材料名称 | 提交要求 |', block['text'])
            self.assertLessEqual(len(block['text']), 220)
        for row in rows:
            self.assertTrue(any(row in b['text'] for b in blocks))

    def test_short_real_values_not_discarded(self):
        blocks = chunking.build_blocks([(1, '一、人数\n1\n二、报名要求\n无')])
        self.assertIn('\n1', '\n'.join(b['text'] for b in blocks))
        self.assertIn('\n无', '\n'.join(b['text'] for b in blocks))

    def test_hard_wrap_without_punctuation_is_bounded(self):
        blocks = chunking.build_blocks([(1, '一、正文\n' + '甲' * 2000)], max_chars=180)
        self.assertTrue(all(len(b['text']) <= 180 for b in blocks))
        self.assertEqual(sum(b['text'].count('甲') for b in blocks), 2000)


if __name__ == '__main__':
    unittest.main()
