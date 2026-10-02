import io
import json
import unittest
import tempfile
from unittest.mock import patch

from PIL import Image
from docx import Document

import document_vision
import document_formats
import parsing
from pypdf import PdfWriter
from pathlib import Path


class VisionTests(unittest.TestCase):
    def test_long_pdf_processed_in_batches(self):
        writer = PdfWriter()
        for _ in range(23): writer.add_blank_page(width=100, height=100)
        raw = io.BytesIO(); writer.write(raw)
        progress = []
        def render(command, **kwargs):
            Path(command[-1] + '.png').write_bytes(b'fixture')
        def recognize(folder, on_result=None):
            results = [{'file': p.name, 'text': '识别内容', 'error': ''} for p in sorted(Path(folder).glob('page-*.png'))]
            self.assertLessEqual(len(results), 6)
            for result in results:
                if on_result: on_result(result)
            return results
        with patch('document_vision.configured', return_value=True), \
             patch('parsing.poppler_binary', return_value='pdftoppm'), \
             patch('parsing.subprocess.run', side_effect=render), \
             patch('document_vision.read_folder', side_effect=recognize):
            pages, meta = parsing.extract(raw.getvalue(), progress=progress.append)
        self.assertEqual(len(pages), 23)
        self.assertTrue(all(p == '识别内容' for p in pages))
        self.assertEqual(meta['unreadable_pages'], [])
        self.assertEqual(progress[-1]['completed'], 23)

    def test_retry_does_not_repeat_successful_images(self):
        from collections import Counter
        counts = Counter()
        def recognize(content):
            counts[content] += 1
            if content == b'failed' and counts[content] == 1:
                raise ValueError('temporary service error')
            return '文字'
        with tempfile.TemporaryDirectory() as folder:
            (Path(folder) / 'page-0001.png').write_bytes(b'ok')
            (Path(folder) / 'page-0002.png').write_bytes(b'failed')
            with patch('document_vision.read_image', side_effect=recognize):
                result = document_vision.read_folder(folder)
        self.assertEqual(counts[b'ok'], 1)
        self.assertEqual(counts[b'failed'], 2)
        self.assertTrue(all(not item['error'] for item in result))

    def response(self, text, finish='stop'):
        return {'choices': [{'finish_reason': finish, 'message': {'content': json.dumps(text)}}]}

    def test_refuses_truncated_result(self):
        with self.assertRaises(ValueError):
            document_vision.validate_response(self.response({'text': '18:', 'uncertain': []}, 'length'))

    def test_refuses_uncertain_numbers(self):
        with self.assertRaises(ValueError):
            document_vision.validate_response(self.response({'text': '报名：9月24日', 'uncertain': ['日期模糊']}))

    def test_uncertain_paragraph_does_not_remove_other_clear_paragraphs(self):
        with self.assertRaises(ValueError) as caught:
            document_vision.validate_response(self.response({
                'text': '每队1—5人\n\n印章：[无法辨认]\n\n发布日期：2026年6月5日',
                'uncertain': ['印章文字']}))
        self.assertEqual(caught.exception.partial_text, '每队1—5人\n\n发布日期：2026年6月5日')

    def test_preserves_verbatim_text(self):
        text = '报名截止：2026年9月24日18:00\n每队1—5人'
        self.assertEqual(document_vision.validate_response(self.response({'text': text, 'uncertain': []})), text)

    def test_refuses_invalid_shape(self):
        for body in ({'text': 123, 'uncertain': []}, {'text': 'x'}, {'text': 'x', 'uncertain': '无'}):
            with self.assertRaises(ValueError):
                document_vision.validate_response(self.response(body))

    def test_structured_blocks_keep_clear_text_and_exclude_uncertain_block(self):
        with self.assertRaises(ValueError) as caught:
            document_vision.validate_response(self.response({'blocks': [
                {'text': '每队1—5人', 'uncertain': False},
                {'text': '印章：不清楚', 'uncertain': True},
                {'text': '2026年6月5日', 'uncertain': False}]}))
        self.assertEqual(caught.exception.partial_text, '每队1—5人\n\n2026年6月5日')

    def test_image_uses_vision_and_does_not_fallback_on_error(self):
        raw = io.BytesIO(); Image.new('RGB', (80, 40)).save(raw, 'PNG')
        with patch('document_vision.configured', return_value=True), \
             patch('document_vision.read_image', side_effect=ValueError('文字不清楚')), \
             patch('document_formats.run_image_ocr') as legacy:
            pages, meta = document_formats.extract_document('x.png', raw.getvalue())
        self.assertEqual(pages, [''])
        self.assertTrue(meta['warning'])
        legacy.assert_not_called()

    def test_word_embedded_image_keeps_position(self):
        raw = io.BytesIO(); Image.new('RGB', (80, 40)).save(raw, 'PNG'); raw.seek(0)
        doc = Document(); doc.add_paragraph('之前'); doc.add_picture(raw); doc.add_paragraph('之后')
        out = io.BytesIO(); doc.save(out)
        with patch('document_vision.configured', return_value=True), \
             patch('document_vision.read_image', return_value='图片里的报名要求'):
            pages, meta = document_formats.extract_document('x.docx', out.getvalue())
        self.assertLess(pages[0].index('之前'), pages[0].index('图片里的报名要求'))
        self.assertLess(pages[0].index('图片里的报名要求'), pages[0].index('之后'))
        self.assertEqual(meta['location_kind'], 'section')

    def test_word_inline_image_keeps_text_on_both_sides(self):
        raw = io.BytesIO(); Image.new('RGB', (80, 40)).save(raw, 'PNG'); raw.seek(0)
        doc = Document(); p = doc.add_paragraph('之前'); p.add_run().add_picture(raw); p.add_run('之后')
        out = io.BytesIO(); doc.save(out)
        with patch('document_vision.configured', return_value=True), \
             patch('document_vision.read_image', return_value='图片内容'):
            pages, _ = document_formats.extract_document('x.docx', out.getvalue())
        self.assertLess(pages[0].index('图片内容'), pages[0].index('之后'))

    def test_pdf_vision_uses_original_page_numbers(self):
        raw = io.BytesIO(); Image.new('RGB', (80, 40)).save(raw, 'PDF')
        with patch('document_vision.configured', return_value=True), \
             patch('parsing.poppler_binary', return_value='pdftoppm'), \
             patch('parsing.subprocess.run'), \
             patch('document_vision.read_folder', return_value=[{'file': 'page-0001.png', 'text': '每队1至5人', 'error': ''}]):
            pages, meta = parsing.extract(raw.getvalue())
        self.assertEqual(pages, ['每队1至5人'])
        self.assertEqual(meta['ocr_pages'], [1])
        self.assertEqual(meta['text_extraction_method'], 'native_text+deepseek_vision')


if __name__ == '__main__':
    unittest.main()
